"""Spike smoke test: drive the PyInstaller bundle and report findings.

Usage: python spike/smoke.py <path to dist/commcare-sync directory>

Reads CORPORA_USERNAME and CORPORA_API_KEY from the environment for the
real export step; skips that step if they are missing.

Every check appends a line to the findings list and the script continues,
so one run yields as much information as possible. The exit code is
non-zero if any check failed.
"""

import os
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.request
from datetime import date, timedelta
from pathlib import Path

import psutil

FINDINGS: list[str] = []
FAILED = False
START_TIMEOUT = 180
TICK_TIMEOUT = 150
STOP_TIMEOUT = 90


def note(message):
    print(f'FINDING: {message}', flush=True)
    FINDINGS.append(message)


def fail(message):
    global FAILED
    FAILED = True
    note(f'FAIL: {message}')


def ok(message):
    note(f'OK: {message}')


def dir_size(path):
    return sum(p.stat().st_size for p in Path(path).rglob('*') if p.is_file())


def free_port():
    import socket

    with socket.socket() as sock:
        sock.bind(('127.0.0.1', 0))
        return sock.getsockname()[1]


def http_status(url):
    request = urllib.request.Request(url, method='GET')
    try:
        with urllib.request.urlopen(request, timeout=5) as response:
            return response.status
    except urllib.error.HTTPError as err:
        return err.code
    except (urllib.error.URLError, ConnectionError, TimeoutError, OSError):
        return None


def run(cmd, **kwargs):
    print('+', ' '.join(map(str, cmd)), flush=True)
    return subprocess.run(
        cmd, capture_output=True, text=True, timeout=600, **kwargs
    )


def check_help(exe):
    result = run([exe, '--run-det', '--help'])
    if result.returncode == 0 and 'commcare-export' in result.stdout:
        ok('--run-det --help exits 0')
    else:
        fail(
            f'--run-det --help exited {result.returncode}:\n'
            f'{result.stdout[-2000:]}\n{result.stderr[-2000:]}'
        )
    result = run([exe, '--run-det', '--version'])
    note(f'--run-det --version: {result.stdout.strip()!r} rc={result.returncode}')


def check_imports(exe):
    result = run([exe, '--spike-imports'])
    print(result.stdout, result.stderr, flush=True)
    for line in result.stdout.splitlines():
        if line.startswith('IMPORT'):
            note(line)
    if result.returncode != 0:
        fail('one or more SQL driver imports failed')
    else:
        ok('all driver imports succeeded')


def write_query_file(path):
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'spike_cases'
    sheet.append(['Data Source', 'Filter Name', 'Filter Value', 'Source Field', 'Field'])
    sheet.append(['case', '', '', 'case_id', 'case_id'])
    sheet.append(['', '', '', 'properties.case_name', 'case_name'])
    sheet.append(['', '', '', 'date_modified', 'date_modified'])
    workbook.save(path)


def check_export(exe, work):
    username = os.environ.get('CORPORA_USERNAME')
    api_key = os.environ.get('CORPORA_API_KEY')
    if not username or not api_key:
        note('SKIP: export (CORPORA_USERNAME/CORPORA_API_KEY not set)')
        return
    query = work / 'query.xlsx'
    write_query_file(query)
    # DET's SQL writer knows Postgres, MySQL, MSSQL and Oracle, not SQLite.
    output = work / 'export.xlsx'
    since = (date.today() - timedelta(days=30)).isoformat()
    result = run(
        [
            exe,
            '--run-det',
            '--query',
            str(query),
            '--project',
            'corpora',
            '--commcare-hq',
            'https://www.commcarehq.org',
            '--username',
            username,
            '--password',
            api_key,
            '--auth-mode',
            'apikey',
            '--since',
            since,
            '--output-format',
            'xlsx',
            '--output',
            str(output),
            '--verbose',
        ],
        cwd=work,
    )
    log = (result.stdout + '\n' + result.stderr).replace(api_key, '***')
    print(log[-6000:], flush=True)
    if result.returncode != 0:
        fail(f'--run-det export exited {result.returncode}')
        return
    try:
        import openpyxl

        workbook = openpyxl.load_workbook(output, read_only=True)
        sheets = workbook.sheetnames
        count = sum(1 for _ in workbook['spike_cases'].iter_rows()) - 1
    except Exception as err:  # noqa: BLE001
        fail(f'export output unreadable: {err!r}')
        return
    ok(f'--run-det export wrote sheets {sheets}; spike_cases has {count} rows')


def descendants(pid):
    try:
        parent = psutil.Process(pid)
    except psutil.NoSuchProcess:
        return []
    return parent.children(recursive=True)


def describe(proc):
    try:
        return f'{proc.pid}:{proc.name()}:{" ".join(proc.cmdline()[1:3])}'
    except psutil.Error:
        return f'{proc.pid}:?'


def check_server(exe, dist, work):
    data_dir = work / 'data'
    data_dir.mkdir()
    port = free_port()
    url = f'http://127.0.0.1:{port}/'
    env = {k: v for k, v in os.environ.items() if not k.startswith('DJANGO')}
    env.pop('COMMCARE_SYNC_DATA_DIR', None)
    started = time.monotonic()
    log_file = open(work / 'server-stdout.log', 'w', encoding='utf-8')
    process = subprocess.Popen(
        [exe, '--server', '--port', str(port), '--data-dir', str(data_dir)],
        stdout=log_file,
        stderr=subprocess.STDOUT,
        env=env,
    )
    try:
        deadline = started + START_TIMEOUT
        status = None
        while time.monotonic() < deadline and process.poll() is None:
            status = http_status(url)
            if status is not None:
                break
            time.sleep(0.5)
        if status in (200, 302):
            note(
                f'cold start to first HTTP response: '
                f'{time.monotonic() - started:.1f}s (status {status})'
            )
        elif process.poll() is not None:
            fail(f'launcher exited early with {process.returncode}')
            return
        else:
            fail(f'no 200/302 from {url} within {START_TIMEOUT}s (got {status})')

        setup_status = http_status(url + 'setup/')
        note(f'/setup/ status: {setup_status}')

        db = data_dir / 'db.sqlite3'
        deadline = time.monotonic() + TICK_TIMEOUT
        tick = None
        while time.monotonic() < deadline and process.poll() is None:
            try:
                with sqlite3.connect(f'file:{db}?mode=ro', uri=True) as conn:
                    conn.execute('pragma busy_timeout=2000')
                    rows = conn.execute(
                        'select success, started, stopped from django_q_task '
                        "where func='apps.schedules.tasks.run_due_schedules'"
                    ).fetchall()
            except sqlite3.Error:
                rows = []
            if rows:
                tick = rows
                break
            time.sleep(2)
        if tick and all(row[0] for row in tick):
            ok(
                f'dispatcher tick succeeded {time.monotonic() - started:.0f}s '
                f'after launch: {tick}'
            )
        elif tick:
            fail(f'dispatcher tick ran but failed: {tick}')
        else:
            fail(f'no dispatcher task recorded within {TICK_TIMEOUT}s')

        try:
            with sqlite3.connect(f'file:{db}?mode=ro', uri=True) as conn:
                schedules = conn.execute(
                    'select func, schedule_type, minutes from django_q_schedule'
                ).fetchall()
            note(f'schedules: {schedules}')
        except sqlite3.Error as err:
            fail(f'could not read django_q_schedule: {err!r}')

        tree = descendants(process.pid)
        note(f'process tree while running: {[describe(p) for p in tree]}')
        if len(tree) < 2:
            fail('expected qcluster plus at least one worker as descendants')
        tracked = [(p.pid, p.create_time()) for p in tree]

        # Stop via the launcher's own shutdown path.
        (data_dir / 'STOP').touch()
        stop_started = time.monotonic()
        try:
            process.wait(timeout=STOP_TIMEOUT)
            note(
                f'launcher exited {process.returncode} '
                f'{time.monotonic() - stop_started:.1f}s after STOP'
            )
        except subprocess.TimeoutExpired:
            fail(f'launcher did not exit within {STOP_TIMEOUT}s of STOP')
            process.kill()
            process.wait()

        time.sleep(2)
        leftovers = []
        for pid, created in tracked:
            try:
                proc = psutil.Process(pid)
                if abs(proc.create_time() - created) < 1:
                    leftovers.append(describe(proc))
            except psutil.NoSuchProcess:
                pass
        dist_path = str(Path(dist).resolve()).lower()
        for proc in psutil.process_iter(['pid', 'exe', 'cmdline']):
            try:
                exe_path = (proc.info['exe'] or '').lower()
            except psutil.Error:
                continue
            if exe_path.startswith(dist_path):
                leftovers.append(describe(proc))
        if leftovers:
            fail(f'processes left behind after stop: {leftovers}')
        else:
            ok('no descendant processes left after stop')
    finally:
        if process.poll() is None:
            process.kill()
            process.wait()
        log_file.close()

    tracebacks = []
    for log in sorted((data_dir / 'logs').glob('*')) + [work / 'server-stdout.log']:
        try:
            text = log.read_text(encoding='utf-8', errors='replace')
        except OSError:
            continue
        if 'Traceback' in text:
            tracebacks.append(log.name)
    if tracebacks:
        fail(f'tracebacks in logs: {tracebacks}')
    else:
        ok('no tracebacks in logs')


def main():
    dist = Path(sys.argv[1])
    # Absolute: the export runs with another working directory.
    exe = (dist / ('commcare-sync.exe' if sys.platform == 'win32' else 'commcare-sync')).resolve()
    if not exe.exists():
        print(f'no executable at {exe}')
        return 2
    note(f'platform: {sys.platform} python: {sys.version.split()[0]}')
    note(f'bundle size: {dir_size(dist) / 1024 / 1024:.0f} MiB')
    work = Path(tempfile.mkdtemp(prefix='spike-'))
    try:
        started = time.monotonic()
        check_help(exe)
        note(f'--run-det --help wall time: {time.monotonic() - started:.1f}s')
        check_imports(exe)
        check_export(exe, work)
        check_server(exe, dist, work)
    finally:
        out_dir = Path(os.environ.get('SPIKE_OUTPUT_DIR', 'spike-output'))
        out_dir.mkdir(exist_ok=True)
        for item in ['server-stdout.log', 'data/logs']:
            source = work / item
            if source.is_dir():
                shutil.copytree(source, out_dir / item.replace('/', '-'), dirs_exist_ok=True)
            elif source.exists():
                shutil.copy(source, out_dir)
        summary = '\n'.join(f'- {line}' for line in FINDINGS)
        (out_dir / 'findings.md').write_text(summary + '\n', encoding='utf-8')
        step_summary = os.environ.get('GITHUB_STEP_SUMMARY')
        if step_summary:
            with open(step_summary, 'a', encoding='utf-8') as fh:
                fh.write(f'## Spike findings ({sys.platform})\n\n{summary}\n')
        print('\n=== FINDINGS ===\n' + summary, flush=True)
    return 1 if FAILED else 0


if __name__ == '__main__':
    sys.exit(main())
